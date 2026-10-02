# dump.py NAME OUT → Folientexte + Notizen als Text (für die Faktenprüfung)
import sys
from pptx import Presentation
p=Presentation(sys.argv[1]+'.pptx'); out=[]
for i,s in enumerate(p.slides,1):
  t=[sh.text_frame.text.replace('\n',' / ') for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
  n=s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''
  out.append(f'=== FOLIE {i} ===\nAUF DER FOLIE: '+' | '.join(t)+'\nNOTIZEN:\n'+n+'\n')
open(sys.argv[2],'w').write('\n'.join(out))
