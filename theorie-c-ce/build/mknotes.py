# mknotes.py NAME → r/NAME/notes.json (Notizen je Folie, für sprechzettel.js)
import sys, json
from pptx import Presentation
n = sys.argv[1]; p = Presentation(n + '.pptx')
json.dump([s.notes_slide.notes_text_frame.text if s.has_notes_slide else '' for s in p.slides], open(f'r/{n}/notes.json', 'w'), ensure_ascii=False)
print(len(p.slides))
