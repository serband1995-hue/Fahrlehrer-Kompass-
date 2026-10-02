# notecheck.py DATEI.pptx → prüft Notizen, Kopfzeile, Klickzahl, Übergänge, verbotene Wörter
from pptx import Presentation
from lxml import etree
import re,sys
f=sys.argv[1]; p=Presentation(f); N=len(p.slides); prob=[]; tot=0
BAD=['KI-generiert','Higgsfield','KI generiert','Gruppe ','Gruppenarbeit']
for i,s in enumerate(p.slides,1):
    t=s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''
    x=etree.tostring(s._element).decode(); clicks=x.count('nodeType="clickEffect"'); tot+=clicks
    if len(t.strip())<40: prob.append(f'F{i}: keine/kaum Notizen'); continue
    if i<N and '➜' not in t: prob.append(f'F{i}: kein ➜')
    if '▶' not in t: prob.append(f'F{i}: kein ▶')
    if '🖱' not in t: prob.append(f'F{i}: kein 🖱')
    m=re.search(r'FOLIE (\d+) VON (\d+)',t)
    if not m or int(m.group(1))!=i or int(m.group(2))!=N: prob.append(f'F{i}: Kopf falsch')
    m2=re.search(r'· (\d+) Klicks?',t.split('\n')[0])
    real=clicks
    if m2 and int(m2.group(1))!=real: prob.append(f'F{i}: Kopf sagt {m2.group(1)} Klicks, echt {real}')
    if not m2 and real: prob.append(f'F{i}: Kopf ohne Klicks, echt {real}')
    # Klickangabe in Notizen prüfen: "Klick 1–N" oder höchste "Klick N"
    nums=[int(a) for a in re.findall(r'Klick (\d+)',t)]+[int(b) for a,b in re.findall(r'Klick (\d+)[–-](\d+)',t)]
    if real and nums and max(nums)!=real and 'Jeder Klick' not in t: prob.append(f'F{i}: Notiz nennt bis Klick {max(nums)}, echt {real}')
    if real==0 and nums: prob.append(f'F{i}: Notiz nennt Klicks, Folie hat keine')
    if real and not nums and 'Jeder Klick' not in t: prob.append(f'F{i}: {real} Klicks, aber keine Klickangabe')
    for b in BAD:
        if b in t: prob.append(f'F{i}: verbotenes Wort "{b}"')
    for sh in s.shapes:
        if sh.has_text_frame and any(b in sh.text_frame.text for b in BAD[:3]): prob.append(f'F{i}: verbotenes Wort auf Folie')
print(f'{f}: {N} Folien, {tot} Klicks, {len(prob)} Auffälligkeiten'); [print('  ',x) for x in prob]
