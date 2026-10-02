import copy, re, io
from pptx import Presentation
from pptx.oxml.ns import qn
A=Presentation('gs910_neu.pptx'); B=Presentation('../ext/x910.pptx')
RNS='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
def clone(src):
    dst=A.slides.add_slide(A.slide_layouts[0])
    # Bilder übernehmen
    rmap={}
    for rId,rel in src.part.rels.items():
        if rel.reltype.endswith('/image'):
            _,new=dst.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
            rmap[rId]=new
    sx=src._element; dx=dst._element
    dx.replace(dx.find(qn('p:cSld')), copy.deepcopy(sx.find(qn('p:cSld'))))
    after=False
    for ch in list(sx):
        if ch.tag==qn('p:clrMapOvr'): after=True; continue
        if after: dx.append(copy.deepcopy(ch))
    for el in dx.iter():
        for att in ('{%s}embed'%RNS,'{%s}link'%RNS,'{%s}id'%RNS):
            v=el.get(att)
            if v is not None:
                if v not in rmap: raise Exception('rel? '+v)
                el.set(att,rmap[v])
    t=src.notes_slide.notes_text_frame.text
    t=re.sub(r'^FOLIE \d+ VON \d+ · [^\n]*\n\n','',t)
    dst.notes_slide.notes_text_frame.text=t
    return dst
orig=list(A.slides._sldIdLst)  # sldId-Elemente 1..77
new=[clone(s) for s in B.slides]
newids=list(A.slides._sldIdLst)[77:]
O=lambda a,b:[orig[i-1] for i in range(a,b+1)]
X=lambda a,b:[newids[i-1] for i in range(a,b+1)]
seq=O(1,48)+X(1,6)+O(49,51)+X(7,8)+O(52,52)+X(9,23)+O(53,53)+O(73,73)+X(24,25)+O(54,68)+X(26,30)+O(69,70)+X(31,40)+O(71,71)+X(41,52)+O(74,77)
drop=orig[71]
lst=A.slides._sldIdLst
for e in list(lst): lst.remove(e)
for e in seq: lst.append(e)
# Folie 72 (alte Sammel-Auflösung) entfernen: Beziehung lösen
A.part.drop_rel(drop.get(qn('r:id')))
print('Folien',len(A.slides))
A.save('gs910_180.pptx')
