from pptx import Presentation
from lxml import etree
import re,sys
D={'1+2':'notes_fix/out/Abend_Lektion01_02.pptx','3+4':'notes_fix/out/Abend_Lektion03_04.pptx','7':'notes_fix/out/Lektion07_Andere_Verkehrsteilnehmer.pptx','8':'notes_fix/out/Lektion08_Geschwindigkeit_Abstand_Umwelt.pptx','11':'notes_fix/out/Lektion11_Besondere_Situationen.pptx','12':'notes_fix/out/Lektion12_Lebenslanges_Lernen.pptx'}
for k,f in D.items():
    p=Presentation(f); N=len(p.slides); prob=[]
    for i,s in enumerate(p.slides,1):
        t=s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''
        x=etree.tostring(s._element).decode(); clicks=x.count('nodeType="clickEffect"')
        if len(t.strip())<15: prob.append(f'F{i}: keine/kaum Notizen'); continue
        if i<N and '➜' not in t: prob.append(f'F{i}: kein ➜')
        m=re.search(r'FOLIE (\d+) VON (\d+)',t)
        if m and (int(m.group(1))!=i or int(m.group(2))!=N): prob.append(f'F{i}: Kopf {m.group(0)}')
        m=re.search(r'(\d+) Klicks?',t.split('\n')[0]) if m else None
        if m and int(m.group(1))!=clicks: prob.append(f'F{i}: Kopf sagt {m.group(1)} Klicks, echt {clicks}')
        for bad in ['Folie 71','Folie 72','Gruppe 4','Live auf der A3','Lektion 8 · Andere','Lektion 8, Andere','Lektion 7 · Geschw','KI-generiert']:
            if bad in t: prob.append(f'F{i}: "{bad}"')
    print(f'== {k} ({N} Folien): {len(prob)} Auffälligkeiten'); [print('  ',x) for x in prob]
