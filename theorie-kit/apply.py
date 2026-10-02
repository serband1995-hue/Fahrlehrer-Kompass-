import json,sys,zipfile,shutil,os
from pptx import Presentation
k,src,dst=sys.argv[1:4]
J=json.load(open(f'notes_fix/{k}.json'))
p=Presentation(src); N=len(p.slides); nu=0; nf=0; miss=[]
for f in J.get('fehler',[]):
    nt=p.slides[int(f['folie'])-1].notes_slide.notes_text_frame; t=nt.text
    if f['alt'] in t: nt.text=t.replace(f['alt'],f['neu']); nf+=1
    else: miss.append(f['folie'])
for n,u in J['uebergaenge'].items():
    n=int(n)
    if n>=N: continue
    nt=p.slides[n-1].notes_slide.notes_text_frame
    if '➜' in nt.text: continue
    u=u.strip()
    if not u.startswith('➜'): u='➜ '+u
    nt.text=nt.text.rstrip()+'\n'+u; nu+=1
p.save(dst)
# Content-Type jpg sicherstellen
zi=zipfile.ZipFile(dst); tmp=dst+'.tmp'; zo=zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED)
for it in zi.infolist():
    d=zi.read(it.filename)
    if it.filename=='[Content_Types].xml' and b'Extension="jpg"' not in d: d=d.replace(b'<Default ',b'<Default Extension="jpg" ContentType="image/jpeg"/><Default ',1)
    zo.writestr(it,d)
zo.close(); zi.close(); shutil.move(tmp,dst)
print(k,'Übergänge',nu,'Korrekturen',nf,'nicht gefunden',miss)
